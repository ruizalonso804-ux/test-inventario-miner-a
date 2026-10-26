"""
generar_datos.py
----------------
Genera bodega_ejemplo.xlsx con datos FICTICIOS de una bodega de repuestos
mineros y, si existe index.html en la misma carpeta, actualiza el bloque de
datos de ejemplo embebido en la página.

Modelo de datos:
  - Cada producto tiene un SKU único (descripción, criticidad, consumo,
    lead time y stock mínimo son atributos del SKU).
  - Un mismo SKU puede estar en varias ubicaciones; cada fila trae el stock
    de esa ubicación. El estado se evalúa con el stock total del SKU.
  - Criticidad A solo para repuestos de bombas, chancador, hidráulica,
    eléctrica o transmisión de equipos críticos. EPP, fijaciones y
    consumibles generales quedan en B o C.
  - costo_unitario en USD (ficticio, por SKU).

Uso:
    pip install pandas openpyxl
    python generar_datos.py

Autor: Alonso Ruiz Fredes · Ingeniería Civil Industrial, UNACH
La semilla es fija (42), por lo que cada ejecución produce el mismo archivo.
"""

from __future__ import annotations

import json
import math
import re
from collections import Counter
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Parámetros de la bodega (deben coincidir con CONFIG en index.html)
# ---------------------------------------------------------------------------
NUM_PASILLOS = 10          # A..J
NUM_COLUMNAS = 24
NUM_NIVELES = 5
PROB_VACIO = 0.10          # ~10 % de posiciones vacías
SEMILLA = 42
DIAS_SOBRESTOCK = 180

CARPETA = Path(__file__).resolve().parent
ARCHIVO_XLSX = CARPETA / "bodega_ejemplo.xlsx"
ARCHIVO_HTML = CARPETA / "index.html"

COLUMNAS = [
    "pasillo", "columna", "nivel", "sku", "descripcion", "criticidad",
    "stock_actual", "stock_minimo", "consumo_mensual", "lead_time_dias",
    "costo_unitario",
]

# ---------------------------------------------------------------------------
# Equipos de la faena. True = equipo crítico (su detención para la planta
# o la mina). Solo repuestos de equipos críticos pueden ser criticidad A.
# ---------------------------------------------------------------------------
EQUIPOS = {
    "Chancador primario": True,
    "Molino SAG": True,
    "Molino de bolas": True,
    "Bomba alimentación ciclones": True,
    "Camión de extracción": True,
    "Pala eléctrica": True,
    "Correa overland CV-01": True,
    "Chancador de pebbles": False,
    "Correa de stock CV-05": False,
    "Bomba sumidero planta": False,
    "Bomba agua de proceso": False,
    "Planta de filtros": False,
    "Cargador frontal": False,
    "Bulldozer": False,
    "Motoniveladora": False,
    "Bodega general": False,
}

# ---------------------------------------------------------------------------
# Sistemas del equipo al que pertenece cada repuesto. Solo los repuestos de
# estos sistemas, en equipos críticos, pueden tener criticidad A.
# EPP, fijaciones y consumibles generales quedan siempre en B o C.
# ---------------------------------------------------------------------------
SISTEMAS_A = {"bombas", "chancador", "hidraulica", "electrica", "transmision"}

# ---------------------------------------------------------------------------
# Catálogo ficticio por familia.
#   participacion: fracción objetivo de posiciones ocupadas
#   items: (descripción base, variantes, equipos, prob. de A, sistema, costo USD (mín, máx))
# ---------------------------------------------------------------------------
FAMILIAS = {
    "rodamientos": dict(
        prefijo="ROD", participacion=0.15, consumo=(0.3, 5), lead=(45, 120),
        pesado=False, ubicaciones=(1, 3),
        items=[
            ("Rodamiento rodillos esféricos",
             ["22228 CC", "22232 CC", "22320 CC", "23044 CC", "23148 CC", "23152 CC", "23260 CA", "24152 CA"],
             ["Molino SAG", "Molino de bolas", "Correa overland CV-01", "Chancador de pebbles", "Correa de stock CV-05"],
             0.8, "transmision", (800, 9500)),
            ("Rodamiento rodillos cónicos", ["32222", "32030", "32228", "HM265049", "HH932132", "M268749"],
             ["Camión de extracción", "Pala eléctrica", "Cargador frontal"], 0.6, "transmision", (300, 2800)),
            ("Rodamiento de bolas", ["6316-2RS", "6318 C3", "6220 C3", "6322", "6224 C3", "6413"],
             ["Bomba alimentación ciclones", "Bomba agua de proceso", "Planta de filtros"], 0.5, "bombas", (60, 600)),
            ("Soporte con rodamiento", ["SNL 520", "SNL 524", "SNL 528", "SD 3148", "SD 3152"],
             ["Correa overland CV-01", "Correa de stock CV-05"], 0.5, "transmision", (900, 4500)),
            ("Rodamiento rodillos cilíndricos", ["NU 2232", "NU 2244", "NJ 2330", "NUP 2324"],
             ["Chancador primario", "Pala eléctrica", "Molino de bolas"], 0.7, "transmision", (1200, 8000)),
        ],
    ),
    "sellos": dict(
        prefijo="SEL", participacion=0.11, consumo=(1, 10), lead=(30, 75),
        pesado=False, ubicaciones=(1, 3),
        items=[
            ("Sello mecánico", ["65 mm", "75 mm", "90 mm", "100 mm", "125 mm"],
             ["Bomba alimentación ciclones", "Bomba agua de proceso", "Bomba sumidero planta"], 0.7, "bombas", (900, 4800)),
            ("Kit sellos cilindro", ["levante tolva", "dirección", "suspensión delantera", "suspensión trasera", "freno"],
             ["Camión de extracción"], 0.6, "hidraulica", (250, 1200)),
            ("Kit sellos cilindro", ["levante balde", "empuje", "apertura balde", "inclinación"],
             ["Pala eléctrica", "Cargador frontal"], 0.5, "hidraulica", (300, 1500)),
            ("Sello laberinto muñón", ["alimentación", "descarga", "piñón"],
             ["Molino SAG", "Molino de bolas"], 0.0, "molienda", (2500, 9000)),
            ("Retén radial", ["95x120x12", "180x220x15", "200x240x20", "240x280x20", "300x340x20", "160x190x15"],
             ["Chancador primario", "Chancador de pebbles", "Correa de stock CV-05"], 0.5, "transmision", (40, 400)),
            ("Empaquetadura prensaestopas", ['1/2"', '5/8"', '3/4"'],
             ["Bomba agua de proceso", "Bomba alimentación ciclones"], 0.4, "bombas", (30, 150)),
        ],
    ),
    "revestimientos_bomba": dict(
        prefijo="RBP", participacion=0.10, consumo=(0.3, 3), lead=(60, 120),
        pesado=True, ubicaciones=(1, 3),
        items=[
            ("Impulsor bomba pulpa", ["8x6", "10x8", "12x10", "14x12", "16x14"],
             ["Bomba alimentación ciclones", "Bomba sumidero planta"], 0.9, "bombas", (6000, 28000)),
            ("Voluta bomba pulpa", ["8x6", "10x8", "12x10", "14x12", "16x14"],
             ["Bomba alimentación ciclones", "Bomba sumidero planta"], 0.8, "bombas", (9000, 35000)),
            ("Placa de garganta bomba pulpa", ["8x6", "10x8", "12x10", "14x12", "16x14"],
             ["Bomba alimentación ciclones", "Bomba sumidero planta"], 0.7, "bombas", (2500, 9000)),
            ("Revestimiento placa bastidor", ["10x8", "12x10", "14x12", "16x14"],
             ["Bomba alimentación ciclones", "Bomba sumidero planta"], 0.6, "bombas", (3000, 12000)),
            ("Anillo expulsor bomba pulpa", ["10x8", "12x10", "14x12", "16x14"],
             ["Bomba alimentación ciclones"], 0.5, "bombas", (1500, 5000)),
        ],
    ),
    "chancador": dict(
        prefijo="CHA", participacion=0.08, consumo=(0.2, 2), lead=(90, 180),
        pesado=True, ubicaciones=(1, 3),
        items=[
            ("Manto", ["superior", "inferior", "completo"], ["Chancador primario", "Chancador de pebbles"],
             0.9, "chancador", (35000, 110000)),
            ("Cóncava", ["fila 1", "fila 2", "fila 3", "fila 4"], ["Chancador primario"], 0.9, "chancador", (18000, 60000)),
            ("Cóncava", ["estándar", "gruesa"], ["Chancador de pebbles"], 0.0, "chancador", (8000, 20000)),
            ("Protector de araña", [""], ["Chancador primario"], 0.7, "chancador", (6000, 15000)),
            ("Buje excéntrico", ["interior", "exterior"], ["Chancador primario", "Chancador de pebbles"],
             0.9, "chancador", (15000, 45000)),
            ("Anillo de desgaste taza", [""], ["Chancador de pebbles"], 0.0, "chancador", (5000, 12000)),
            ("Cilindro hidráulico de ajuste", ["completo", "kit reparación"], ["Chancador primario"],
             0.7, "hidraulica", (8000, 30000)),
            ("Placa de alimentación", ["izquierda", "derecha"], ["Chancador primario", "Chancador de pebbles"],
             0.5, "chancador", (1500, 6000)),
            ("Sello de polvo", [""], ["Chancador primario", "Chancador de pebbles"], 0.6, "chancador", (400, 1500)),
        ],
    ),
    "mangueras": dict(
        prefijo="MAN", participacion=0.10, consumo=(2, 18), lead=(10, 30),
        pesado=False, ubicaciones=(1, 3),
        items=[
            ("Manguera hidráulica 4SH",
             ['3/4" x 1 m', '1" x 1,5 m', '1" x 2,5 m', '1 1/4" x 2 m', '1 1/2" x 2 m', '2" x 3 m'],
             ["Camión de extracción", "Pala eléctrica", "Cargador frontal"], 0.4, "hidraulica", (90, 600)),
            ("Manguera alta presión 2SN", ['3/8" x 0,6 m', '1/2" x 0,8 m', '1/2" x 1,5 m', '5/8" x 1 m', '3/4" x 1,2 m'],
             ["Camión de extracción", "Bulldozer", "Motoniveladora"], 0.3, "hidraulica", (40, 250)),
            ("Manguera succión pulpa", ['6" x 3 m', '8" x 3 m', '10" x 3 m'],
             ["Bomba alimentación ciclones", "Bomba sumidero planta"], 0.4, "bombas", (600, 2200)),
            ("Kit mangueras cilindro", ["levante tolva", "dirección"], ["Camión de extracción"],
             0.6, "hidraulica", (700, 2400)),
        ],
    ),
    "correas": dict(
        prefijo="COR", participacion=0.08, consumo=(1, 12), lead=(20, 60),
        pesado=False, ubicaciones=(1, 3),
        items=[
            ("Correa en V", ["B-86", "B-112", "C-120", "C-158", "5V-1500", "8V-2000"],
             ["Bomba agua de proceso", "Planta de filtros", "Chancador de pebbles"], 0.0, "transmision", (25, 180)),
            ("Polín de carga CEMA D", ['6" x 1.200 mm', '6" x 1.400 mm', '7" x 1.400 mm', '7" x 1.600 mm'],
             ["Correa overland CV-01", "Correa de stock CV-05"], 0.0, "consumible", (90, 450)),
            ("Polín de retorno", ['6" x 1.200 mm', '6" x 1.400 mm'], ["Correa overland CV-01", "Correa de stock CV-05"],
             0.0, "consumible", (70, 300)),
            ("Polín de impacto", ['6" x 1.200 mm', '7" x 1.400 mm'], ["Correa overland CV-01", "Correa de stock CV-05"],
             0.0, "consumible", (120, 500)),
            ("Kit empalme vulcanizado", ["EP-800", "ST-1600"], ["Correa overland CV-01", "Correa de stock CV-05"],
             0.0, "consumible", (2500, 9000)),
            ("Raspador primario", ["1.200 mm", "1.400 mm", "1.600 mm"], ["Correa overland CV-01", "Correa de stock CV-05"],
             0.0, "consumible", (400, 1600)),
        ],
    ),
    "filtros": dict(
        prefijo="FIL", participacion=0.10, consumo=(6, 50), lead=(15, 40),
        pesado=False, ubicaciones=(1, 3),
        items=[
            ("Filtro aceite motor", [""], ["Camión de extracción", "Cargador frontal", "Bulldozer", "Motoniveladora"],
             0.0, "consumible", (40, 120)),
            ("Filtro aire primario", [""], ["Camión de extracción", "Cargador frontal", "Bulldozer", "Motoniveladora"],
             0.0, "consumible", (120, 450)),
            ("Filtro aire secundario", [""], ["Camión de extracción", "Cargador frontal", "Bulldozer"],
             0.0, "consumible", (60, 200)),
            ("Filtro hidráulico retorno 10 µm", [""], ["Pala eléctrica", "Camión de extracción", "Molino SAG", "Cargador frontal"],
             0.0, "consumible", (80, 350)),
            ("Filtro combustible separador de agua", [""], ["Camión de extracción", "Bulldozer", "Motoniveladora", "Cargador frontal"],
             0.0, "consumible", (50, 180)),
            ("Filtro aceite lubricación", [""], ["Molino SAG", "Molino de bolas", "Chancador primario", "Chancador de pebbles"],
             0.0, "consumible", (150, 600)),
            ("Filtro transmisión", [""], ["Camión de extracción", "Cargador frontal", "Bulldozer"], 0.0, "consumible", (60, 220)),
            ("Filtro refrigerante", [""], ["Camión de extracción", "Cargador frontal"], 0.0, "consumible", (30, 90)),
            ("Filtro respiradero estanque hidráulico", [""], ["Pala eléctrica", "Camión de extracción"],
             0.0, "consumible", (25, 80)),
        ],
    ),
    "camion": dict(
        prefijo="CAM", participacion=0.09, consumo=(0.5, 6), lead=(30, 90),
        pesado=False, ubicaciones=(1, 3),
        items=[
            ("Disco de freno húmedo", ["delantero", "trasero"], ["Camión de extracción"], 0.7, "transmision", (1800, 6500)),
            ("Kit reparación cilindro suspensión", ["delantero", "trasero"], ["Camión de extracción"], 0.7, "hidraulica", (2500, 9000)),
            ("Bomba hidráulica", ["de levante", "de dirección", "de frenos"], ["Camión de extracción"], 0.8, "hidraulica", (9000, 32000)),
            ("Turbocompresor motor", ["banco izquierdo", "banco derecho"], ["Camión de extracción"], 0.0, "motor", (5000, 14000)),
            ("Inyector combustible", [""], ["Camión de extracción"], 0.0, "motor", (900, 2600)),
            ("Bomba de agua motor", [""], ["Camión de extracción"], 0.6, "bombas", (1500, 4500)),
            ("Motor de arranque 24 V", [""], ["Camión de extracción"], 0.5, "electrica", (1200, 3500)),
            ("Alternador 24 V 140 A", [""], ["Camión de extracción"], 0.4, "electrica", (1500, 4200)),
            ("Sensor velocidad rueda", ["izquierda", "derecha"], ["Camión de extracción"], 0.3, "electrica", (250, 900)),
            ("Acumulador de freno", [""], ["Camión de extracción"], 0.6, "hidraulica", (1800, 6000)),
            ("Cilindro de dirección", [""], ["Camión de extracción"], 0.7, "hidraulica", (4500, 14000)),
            ("Kit pasadores tolva", [""], ["Camión de extracción"], 0.0, "consumible", (600, 2000)),
            ("Kit empaquetaduras culata", [""], ["Camión de extracción"], 0.0, "motor", (400, 1500)),
            ("Faro LED 24 V", [""], ["Camión de extracción"], 0.0, "consumible", (180, 600)),
        ],
    ),
    "electricos": dict(
        prefijo="ELE", participacion=0.08, consumo=(0.3, 6), lead=(30, 120),
        pesado=False, ubicaciones=(1, 2),
        items=[
            ("Motor eléctrico", ["75 kW", "160 kW", "250 kW"],
             ["Bomba alimentación ciclones", "Bomba agua de proceso", "Correa de stock CV-05"], 0.8, "electrica", (9000, 45000)),
            ("Variador de frecuencia", ["90 kW", "200 kW"], ["Bomba alimentación ciclones", "Correa overland CV-01"],
             0.8, "electrica", (12000, 38000)),
            ("Contactor tripolar", ["95 A", "185 A", "400 A"], ["Bomba agua de proceso", "Planta de filtros", "Correa de stock CV-05"],
             0.3, "electrica", (150, 1800)),
            ("Relé de protección de motor", [""], ["Molino SAG", "Chancador primario", "Bomba alimentación ciclones"],
             0.6, "electrica", (900, 4500)),
            ("Sensor de vibración", ["triaxial", "acelerómetro"], ["Molino SAG", "Molino de bolas", "Chancador primario"],
             0.5, "electrica", (400, 2500)),
            ("Escobillas motor", ["levante", "empuje", "giro"], ["Pala eléctrica"], 0.6, "electrica", (200, 900)),
            ("Interruptor de tirón", [""], ["Correa overland CV-01", "Correa de stock CV-05"], 0.4, "electrica", (250, 900)),
            ("Fusible NH", ["160 A", "250 A", "400 A"], ["Bodega general"], 0.0, "consumible", (15, 90)),
            ("Luminaria LED uso minero", ["150 W", "240 W"], ["Bodega general"], 0.0, "consumible", (250, 800)),
        ],
    ),
    "epp": dict(
        prefijo="EPP", participacion=0.06, consumo=(40, 300), lead=(7, 20),
        pesado=False, ubicaciones=(2, 4),
        items=[
            ("Guante nitrilo anticorte", ["talla 8", "talla 9", "talla 10"], ["Bodega general"], 0.0, "epp", (3, 9)),
            ("Lente de seguridad", ["claro", "oscuro"], ["Bodega general"], 0.0, "epp", (2, 8)),
            ("Respirador medio rostro", ["talla M", "talla L"], ["Bodega general"], 0.0, "epp", (25, 60)),
            ("Filtro respirador P100 (par)", [""], ["Bodega general"], 0.0, "epp", (12, 30)),
            ("Tapón auditivo reutilizable", [""], ["Bodega general"], 0.0, "epp", (1, 4)),
            ("Casco de seguridad con barbiquejo", ["blanco", "amarillo"], ["Bodega general"], 0.0, "epp", (18, 45)),
            ("Chaleco reflectante clase 2", ["talla L", "talla XL"], ["Bodega general"], 0.0, "epp", (8, 20)),
        ],
    ),
    "fijacion": dict(
        prefijo="FIJ", participacion=0.05, consumo=(30, 250), lead=(10, 30),
        pesado=False, ubicaciones=(2, 4),
        items=[
            ("Perno grado 8", ['3/4" x 4"', '1" x 6"', '1 1/4" x 8"'], ["Bodega general"], 0.0, "fijacion", (2, 12)),
            ("Perno hexagonal clase 10.9", ["M20 x 100", "M24 x 120", "M30 x 150"], ["Bodega general"], 0.0, "fijacion", (3, 15)),
            ("Tuerca autofrenante", ["M24", "M30"], ["Bodega general"], 0.0, "fijacion", (0.8, 4)),
            ("Golilla presión", ["M24", "M30"], ["Bodega general"], 0.0, "fijacion", (0.2, 1.5)),
        ],
    ),
}

# Familias preferidas por pasillo (coherencia espacial del mapa)
FAMILIAS_POR_PASILLO = [
    ["epp", "fijacion", "filtros"],          # A: cerca de despacho, alta rotación
    ["filtros", "mangueras"],                # B
    ["mangueras", "correas"],                # C
    ["correas", "sellos"],                   # D
    ["sellos", "rodamientos"],               # E
    ["rodamientos"],                         # F
    ["rodamientos", "electricos"],           # G
    ["camion", "electricos"],                # H
    ["revestimientos_bomba", "chancador"],   # I: piezas pesadas
    ["chancador", "revestimientos_bomba"],   # J
]

# Probabilidad del estado objetivo según criticidad: el riesgo se concentra en A y B
PROB_ESTADO = {
    "A": {"quiebre": 0.10, "critico": 0.24, "bajo_minimo": 0.08, "sobrestock": 0.04, "normal": 0.54},
    "B": {"quiebre": 0.07, "critico": 0.00, "bajo_minimo": 0.20, "sobrestock": 0.07, "normal": 0.66},
    "C": {"quiebre": 0.01, "critico": 0.00, "bajo_minimo": 0.06, "sobrestock": 0.16, "normal": 0.77},
}


def letra_pasillo(i: int) -> str:
    return chr(ord("A") + i)


def calcular_estado(stock, minimo, consumo, lead, crit) -> str:
    """Misma lógica que la página web (evaluada con el stock total del SKU)."""
    if stock <= 0:
        return "quiebre"
    cobertura = stock / (consumo / 30) if consumo > 0 else math.inf
    if crit == "A" and cobertura < lead:
        return "critico"
    if stock < minimo:
        return "bajo_minimo"
    if cobertura > DIAS_SOBRESTOCK:
        return "sobrestock"
    return "normal"


# ---------------------------------------------------------------------------
# Catálogo de SKU
# ---------------------------------------------------------------------------
def construir_catalogo(rng) -> dict[str, list[dict]]:
    """Crea los productos posibles de cada familia (en orden aleatorio)."""
    catalogo = {}
    for fam, f in FAMILIAS.items():
        productos = []
        for base, variantes, equipos, prob_a, sistema, costo in f["items"]:
            for variante in variantes:
                for equipo in equipos:
                    desc = f"{base} {variante}".strip()
                    if equipo != "Bodega general":
                        desc = f"{desc} · {equipo}"
                    if sistema in SISTEMAS_A and EQUIPOS[equipo]:
                        crit = "A" if rng.random() < prob_a else "B"
                    elif sistema in ("epp", "fijacion"):
                        crit = "B" if rng.random() < 0.15 else "C"
                    elif EQUIPOS[equipo]:
                        crit = "B" if rng.random() < 0.6 else "C"
                    else:
                        crit = "B" if rng.random() < 0.4 else "C"
                    productos.append(dict(familia=fam, descripcion=desc, criticidad=crit,
                                          sistema=sistema, rango_costo=costo))
        orden = rng.permutation(len(productos))
        catalogo[fam] = [productos[i] for i in orden]
    return catalogo


def completar_atributos(rng, prod: dict, numero: int) -> dict:
    f = FAMILIAS[prod["familia"]]
    lo, hi = f["consumo"]
    consumo = round(float(rng.uniform(lo, hi)), 1) if hi <= 6 else float(rng.integers(lo, hi + 1))
    lead = int(rng.integers(f["lead"][0], f["lead"][1] + 1))
    minimo = max(1, math.ceil(consumo / 30 * lead * rng.uniform(0.9, 1.3)))
    # Costo log-uniforme dentro del rango del ítem, redondeado según magnitud
    lo_c, hi_c = prod["rango_costo"]
    costo = float(math.exp(rng.uniform(math.log(lo_c), math.log(hi_c))))
    costo = round(costo, 2) if costo < 10 else round(costo) if costo < 1000 else round(costo, -1)
    return dict(prod, sku=f"{f['prefijo']}-{numero:05d}", consumo_mensual=consumo,
                lead_time_dias=lead, stock_minimo=minimo, costo_unitario=costo, ubicaciones=[])


def stock_total_objetivo(rng, sku: dict) -> int:
    """Stock total del SKU coherente con un estado objetivo según criticidad."""
    probs = PROB_ESTADO[sku["criticidad"]]
    objetivo = rng.choice(list(probs), p=list(probs.values()))
    diario = sku["consumo_mensual"] / 30
    lead, minimo = sku["lead_time_dias"], sku["stock_minimo"]
    cob_lead = math.ceil(diario * lead)
    tope_normal = math.floor(diario * DIAS_SOBRESTOCK)

    if objetivo == "quiebre":
        return 0
    if objetivo == "critico":
        return max(1, math.floor(diario * lead * rng.uniform(0.15, 0.9)))
    if objetivo == "bajo_minimo":
        piso = cob_lead if sku["criticidad"] == "A" else 1
        return int(rng.integers(piso, minimo)) if piso < minimo else max(1, minimo - 1)
    if objetivo == "sobrestock":
        return math.ceil(diario * rng.uniform(200, 420))
    piso = max(minimo, cob_lead if sku["criticidad"] == "A" else 0)
    return int(rng.integers(piso, max(piso + 1, tope_normal)))


def repartir(rng, total: int, n: int) -> list[int]:
    """Reparte el stock total entre n ubicaciones (≥1 cada una si alcanza)."""
    if total <= 0:
        return [0] * n
    if total < n:
        return [1] * total + [0] * (n - total)
    pesos = rng.dirichlet(np.ones(n) * 2)
    partes = [1 + int(x) for x in np.floor(pesos * (total - n))]
    partes[0] += total - sum(partes)
    return partes


# ---------------------------------------------------------------------------
# Ubicación en la bodega
# ---------------------------------------------------------------------------
def generar() -> tuple[pd.DataFrame, list[dict]]:
    rng = np.random.default_rng(SEMILLA)
    catalogo = construir_catalogo(rng)

    # Posiciones ocupadas (≈10 % queda vacía)
    libres = {p: [] for p in range(NUM_PASILLOS)}
    for p in range(NUM_PASILLOS):
        for c in range(1, NUM_COLUMNAS + 1):
            for n in range(1, NUM_NIVELES + 1):
                if rng.random() >= PROB_VACIO:
                    libres[p].append((c, n))

    pasillos_de = {fam: [p for p in range(NUM_PASILLOS)
                         if fam in FAMILIAS_POR_PASILLO[p % len(FAMILIAS_POR_PASILLO)]]
                   for fam in FAMILIAS}
    familias = list(FAMILIAS)
    participacion = np.array([FAMILIAS[f]["participacion"] for f in familias])
    participacion /= participacion.sum()

    skus: list[dict] = []
    por_familia: dict[str, list[dict]] = {f: [] for f in familias}

    def permitido(fam, pos):
        return not FAMILIAS[fam]["pesado"] or pos[1] <= 2   # lo pesado en niveles 1-2

    while any(libres.values()):
        fam = familias[rng.choice(len(familias), p=participacion)]
        # Pasillo: 80 % uno preferido de la familia, si tiene espacio
        candidatos = [p for p in pasillos_de[fam] if any(permitido(fam, x) for x in libres[p])]
        if not candidatos or rng.random() >= 0.8:
            candidatos = [p for p in libres if any(permitido(fam, x) for x in libres[p])]
        if not candidatos:
            continue
        p = candidatos[rng.integers(len(candidatos))]

        if catalogo[fam]:
            sku = completar_atributos(rng, catalogo[fam].pop(), len(skus) + 10001)
            skus.append(sku)
            por_familia[fam].append(sku)
            lo, hi = FAMILIAS[fam]["ubicaciones"]
            cantidad = int(rng.integers(lo, hi + 1))
        else:   # catálogo agotado: una ubicación más para el SKU con menos ubicaciones
            sku = min(por_familia[fam], key=lambda x: len(x["ubicaciones"]))
            cantidad = 1

        # Ubicaciones contiguas alrededor de un punto al azar del pasillo
        opciones = [x for x in libres[p] if permitido(fam, x)]
        inicio = opciones[rng.integers(len(opciones))]
        opciones.sort(key=lambda x: abs(x[0] - inicio[0]) + 0.5 * abs(x[1] - inicio[1]))
        for pos in opciones[:cantidad]:
            libres[p].remove(pos)
            sku["ubicaciones"].append((p, *pos))

    # Stock por SKU y reparto entre sus ubicaciones
    filas = []
    for sku in skus:
        sku["stock_total"] = stock_total_objetivo(rng, sku)
        partes = repartir(rng, sku["stock_total"], len(sku["ubicaciones"]))
        for (p, c, n), stock in zip(sorted(sku["ubicaciones"]), partes):
            filas.append({
                "pasillo": letra_pasillo(p), "columna": c, "nivel": n,
                "sku": sku["sku"], "descripcion": sku["descripcion"],
                "criticidad": sku["criticidad"], "stock_actual": int(stock),
                "stock_minimo": int(sku["stock_minimo"]),
                "consumo_mensual": sku["consumo_mensual"],
                "lead_time_dias": sku["lead_time_dias"],
                "costo_unitario": sku["costo_unitario"],
            })
    df = pd.DataFrame(filas, columns=COLUMNAS).sort_values(["pasillo", "columna", "nivel"])
    return df.reset_index(drop=True), skus


# ---------------------------------------------------------------------------
# Salidas
# ---------------------------------------------------------------------------
def escribir_excel(df: pd.DataFrame) -> None:
    leeme = pd.DataFrame(
        [
            ("pasillo", "Letra del rack/pasillo (A..J)"),
            ("columna", f"Columna del rack (1..{NUM_COLUMNAS})"),
            ("nivel", f"Nivel en altura (1..{NUM_NIVELES}, 1 = piso)"),
            ("sku", "Código único del repuesto (puede repetirse en varias ubicaciones)"),
            ("descripcion", "Descripción del repuesto y equipo al que pertenece"),
            ("criticidad", "A (bombas, chancador, hidráulica, eléctrica o transmisión de equipo crítico), B (media) o C (baja)"),
            ("stock_actual", "Unidades en ESTA ubicación"),
            ("stock_minimo", "Punto mínimo de reposición del SKU"),
            ("consumo_mensual", "Consumo promedio mensual del SKU (acepta decimales con coma o punto)"),
            ("lead_time_dias", "Días de reposición desde la orden de compra"),
            ("costo_unitario", "Costo unitario del SKU en USD (opcional; activa los indicadores de valor)"),
            ("", ""),
            ("Nota", "DATOS DE EJEMPLO, ficticios. Posiciones sin fila = vacías."),
            ("Nota", "El estado se calcula con el stock total del SKU (suma de sus ubicaciones)."),
            ("Autor", "Alonso Ruiz Fredes · Ingeniería Civil Industrial, UNACH"),
        ],
        columns=["campo", "descripcion"],
    )

    with pd.ExcelWriter(ARCHIVO_XLSX, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="stock", index=False)
        leeme.to_excel(xw, sheet_name="leeme", index=False)

        encabezado = PatternFill("solid", fgColor="1F2937")
        for nombre, anchos in (
            ("stock", [9, 9, 7, 11, 58, 11, 13, 13, 16, 15, 15]),
            ("leeme", [18, 75]),
        ):
            ws = xw.sheets[nombre]
            ws.freeze_panes = "A2"
            for i, ancho in enumerate(anchos, start=1):
                ws.column_dimensions[get_column_letter(i)].width = ancho
            for celda in ws[1]:
                celda.font = Font(bold=True, color="FFFFFF")
                celda.fill = encabezado
                celda.alignment = Alignment(horizontal="center")
        xw.sheets["stock"].auto_filter.ref = xw.sheets["stock"].dimensions


def embeber_en_html(df: pd.DataFrame) -> bool:
    """Reemplaza el JSON del bloque <script id="datos-ejemplo"> en index.html."""
    if not ARCHIVO_HTML.exists():
        return False
    datos = {"columnas": COLUMNAS, "filas": df.values.tolist()}
    texto_json = json.dumps(datos, ensure_ascii=False, separators=(",", ":"))
    html = ARCHIVO_HTML.read_text(encoding="utf-8")
    patron = re.compile(
        r'(<script id="datos-ejemplo" type="application/json">)(.*?)(</script>)',
        re.S,
    )
    if not patron.search(html):
        return False
    html = patron.sub(lambda m: m.group(1) + texto_json + m.group(3), html, count=1)
    ARCHIVO_HTML.write_text(html, encoding="utf-8")
    return True


def main() -> None:
    df, skus = generar()
    escribir_excel(df)

    total = NUM_PASILLOS * NUM_COLUMNAS * NUM_NIVELES
    print(f"Archivo generado: {ARCHIVO_XLSX.name}")
    print(f"Posiciones: {total} | ocupadas: {len(df)} | vacías: {total - len(df)} | SKU: {len(skus)}")

    familia_de = {s["sku"]: s["familia"] for s in skus}
    fam = df.sku.map(familia_de).value_counts(normalize=True)
    print("Participación por familia (posiciones):")
    for f, v in fam.items():
        print(f"  {f:<22} {v:6.1%}")
    print(f"  EPP + fijación         {fam.get('epp', 0) + fam.get('fijacion', 0):6.1%}")

    estado_sku = {s["sku"]: calcular_estado(s["stock_total"], s["stock_minimo"], s["consumo_mensual"],
                                            s["lead_time_dias"], s["criticidad"]) for s in skus}
    tabla = pd.crosstab(df.sku.map(estado_sku), df.criticidad, margins=True, margins_name="total")
    print("Posiciones por estado y criticidad:")
    print(tabla.to_string())
    a_fuera = [s for s in skus if s["criticidad"] == "A"
               and (s["sistema"] not in SISTEMAS_A
                    or not any(EQUIPOS[e] for e in EQUIPOS if s["descripcion"].endswith(e)))]
    print(f"SKU A fuera de sistemas/equipos permitidos: {len(a_fuera)} (debe ser 0)")
    print("Criticidad por sistema (SKU):")
    print(pd.crosstab(pd.Series([s["sistema"] for s in skus], name="sistema"),
                      pd.Series([s["criticidad"] for s in skus], name="crit")).to_string())
    print(f"SKU por cantidad de ubicaciones: {dict(sorted(Counter(len(s['ubicaciones']) for s in skus).items()))}")

    if embeber_en_html(df):
        print("Datos de ejemplo actualizados dentro de index.html")


if __name__ == "__main__":
    main()
